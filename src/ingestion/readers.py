"""Readers for the three source systems. Read-only: no reader ever mutates a source.

Each returns the rawest faithful DataFrame — cleaning belongs to
src/reconciliation, so data-quality problems stay visible and countable
instead of being fixed silently at the door.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import sqlalchemy as sa
from openpyxl import load_workbook

from .. import config


def read_hr(path: Path | None = None) -> pd.DataFrame:
    """HR export from the Employees sheet of hr_system.xlsx (openpyxl, values only)."""
    path = path or config.HR_XLSX
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Employees"]
    rows = ws.iter_rows(values_only=True)
    header = [str(h) for h in next(rows)]
    df = pd.DataFrame(rows, columns=header)
    wb.close()
    return df


def read_financials(path: Path | None = None) -> pd.DataFrame:
    # dtype=object: amounts may be numbers OR currency strings — do not let
    # pandas guess and coerce; normalization handles it explicitly.
    return pd.read_csv(path or config.FINANCIALS_CSV, dtype=object)


def read_headcount_plan(path: Path | None = None) -> pd.DataFrame:
    """Workforce planning template (planned FTE per department per month)."""
    return pd.read_csv(path or config.HEADCOUNT_PLAN_CSV, dtype={"department": str})


def read_cost_centre_map(path: Path | None = None) -> pd.DataFrame:
    return pd.read_csv(path or config.COST_CENTRE_MAP_CSV, dtype=str)


def read_contractors(db_url: str | None = None) -> pd.DataFrame:
    engine = sa.create_engine(db_url or config.contractor_db_url())
    try:
        with engine.connect() as conn:
            # sa.text with a fixed statement — no string interpolation anywhere.
            return pd.read_sql(sa.text("SELECT * FROM contractor_roster"), conn)
    finally:
        engine.dispose()
