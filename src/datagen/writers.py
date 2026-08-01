"""Write the three source systems to their native formats.

- HR      -> data/sources/hr_system.xlsx  (openpyxl; Summary sheet has live
             SUMIFS / COUNTIFS / XLOOKUP formulas so the workbook genuinely
             exercises Excel skills, not just a CSV renamed .xlsx)
- Finance -> data/sources/financials.csv + data/lookup/cost_centre_map.csv
- Contractors -> `contractor_roster` table via SQLAlchemy (SQLite by default,
             Postgres if HELIOS_DB_URL is set). Dates stored as TEXT on
             purpose — the mixed formats ARE the data-quality defect.

Note on PivotTables: openpyxl cannot create a PivotTable from scratch (it can
only preserve existing ones), so the Summary sheet reproduces the same
aggregation with SUMIFS/COUNTIFS formulas. Documented in docs/data_dictionary.md.
"""

from __future__ import annotations

import pandas as pd
import sqlalchemy as sa
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils.dataframe import dataframe_to_rows

from .. import config
from . import org_model as om


def write_hr_xlsx(df: pd.DataFrame, path=None) -> None:
    path = path or config.HR_XLSX
    wb = Workbook()
    ws = wb.active
    ws.title = "Employees"
    for row in dataframe_to_rows(df, index=False, header=True):
        ws.append(row)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    n = len(df) + 1  # data extends to row n
    s = wb.create_sheet("Summary")
    s["A1"] = "Headcount by department (SUMIFS/COUNTIFS over Employees)"
    s["A1"].font = Font(bold=True)
    s["A2"], s["B2"], s["C2"] = "Department", "Active rows", "Sum FTE (raw, mixed scales)"
    canonical = [d.name for d in om.DEPARTMENTS]
    for i, dept in enumerate(canonical, start=3):
        s[f"A{i}"] = dept
        # Counts rows whose department starts with the canonical name and has
        # no termination date — deliberately naive, as a real analyst's first
        # pass would be; reconciliation does it properly.
        s[f"B{i}"] = (
            f'=COUNTIFS(Employees!C2:C{n},A{i}&"*",Employees!F2:F{n},"")'
        )
        s[f"C{i}"] = f'=SUMIFS(Employees!G2:G{n},Employees!C2:C{n},A{i}&"*")'
    r = len(canonical) + 4
    s[f"A{r}"] = "Employee lookup (XLOOKUP)"
    s[f"A{r}"].font = Font(bold=True)
    s[f"A{r+1}"] = "employee_id"
    s[f"B{r+1}"] = df["employee_id"].iloc[0]
    s[f"A{r+2}"] = "full_name"
    s[f"B{r+2}"] = (
        f"=XLOOKUP(B{r+1},Employees!A2:A{n},Employees!B2:B{n})"
    )
    s[f"A{r+3}"] = "department"
    s[f"B{r+3}"] = (
        f"=XLOOKUP(B{r+1},Employees!A2:A{n},Employees!C2:C{n})"
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_financials_csv(df: pd.DataFrame, path=None) -> None:
    path = path or config.FINANCIALS_CSV
    path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(path, index=False)


def write_headcount_plan_csv(truth, path=None) -> None:
    """The workforce planning template — the one source an analyst maintains
    by hand, so it ships clean (canonical names, ISO months, no injected mess)."""
    path = path or config.HEADCOUNT_PLAN_CSV
    path.parent.mkdir(parents=True, exist_ok=True)
    truth.headcount_plan[["department", "month", "planned_fte"]].to_csv(path, index=False)


def write_lookup_csv(df: pd.DataFrame, path=None) -> None:
    path = path or config.COST_CENTRE_MAP_CSV
    path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(path, index=False)


def write_contractor_table(df: pd.DataFrame, db_url: str | None = None) -> None:
    engine = sa.create_engine(db_url or config.contractor_db_url())
    with engine.begin() as conn:
        df.to_sql("contractor_roster", conn, if_exists="replace", index=False)
    engine.dispose()


def write_truth(truth, out_dir=None) -> None:
    """Ground truth kept separate from sources — for model evaluation only.

    Downstream code must never read these files during reconciliation or
    forecasting; they exist so accuracy can be reported against the real
    generating process (spec §2.3 'known ground truth').
    """
    out = out_dir or config.TRUTH_DIR
    out.mkdir(parents=True, exist_ok=True)
    truth.monthly.to_csv(out / "monthly_truth.csv", index=False)
    truth.budget.to_csv(out / "budget_truth.csv", index=False)
